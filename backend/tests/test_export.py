"""Excel export filters, Google Sheet link / last-synced label and Sync now."""

from __future__ import annotations

from io import BytesIO
from unittest.mock import AsyncMock

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import export
from app.api.main import api_router
from app.auth import security
from app.channels import sheets
from app.channels.sheets import HEADER, SyncReport
from app.config import settings
from app.db import get_session
from app.db.models import (
    AuthProvider,
    ChannelType,
    Company,
    Contact,
    Outreach,
    OutreachState,
    Post,
    Setting,
    User,
    UserRole,
)

CSRF = {security.CSRF_HEADER: security.CSRF_VALUE}


@pytest.fixture
def app(db):
    application = FastAPI()
    application.include_router(api_router)
    application.dependency_overrides[get_session] = lambda: db
    return application


def client_as(app, db, role: UserRole | None) -> TestClient:
    http = TestClient(app, follow_redirects=False, headers=CSRF)
    if role is not None:
        user = User(email=f"{role.value}@example.test", name=role.value, provider=AuthProvider.PASSWORD, role=role)
        db.add(user)
        db.commit()
        token, _ = security.create_session_token(user)
        http.cookies.set(security.SESSION_COOKIE, token)
    return http


@pytest.fixture
def leads(db):
    acme = Company(name="Acme")
    beta = Company(name="Beta Labs")
    db.add_all([acme, beta])
    db.flush()
    post = Post(url="https://linkedin.com/posts/1", company_id=acme.id)
    db.add(post)
    db.flush()
    jane = Contact(company_id=acme.id, name="Jane Doe", email="jane@acme.test")
    raj = Contact(company_id=beta.id, name="Raj Patel", email="raj@beta.test")
    db.add_all([jane, raj])
    db.flush()
    db.add_all(
        [
            Outreach(contact_id=jane.id, post_id=post.id, channel=ChannelType.LINKEDIN, state=OutreachState.DRAFTED),
            Outreach(contact_id=raj.id, channel=ChannelType.EMAIL, state=OutreachState.REPLIED),
        ]
    )
    db.commit()


def names(response) -> list[str]:
    sheet = load_workbook(BytesIO(response.content)).active
    assert [c.value for c in sheet[1]] == HEADER
    return [row[2].value for row in sheet.iter_rows(min_row=2)]


class TestExportFilters:
    def test_needs_a_login(self, app, db, leads):
        assert client_as(app, db, None).get("/api/export.xlsx").status_code == 401

    def test_pending_user_is_refused(self, app, db, leads):
        assert client_as(app, db, UserRole.PENDING).get("/api/export.xlsx").status_code == 403

    def test_viewer_can_download_everything(self, app, db, leads):
        r = client_as(app, db, UserRole.VIEWER).get("/api/export.xlsx")
        assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
        assert names(r) == ["Jane Doe", "Raj Patel"]

    @pytest.mark.parametrize(
        ("params", "expected"),
        [
            ({"status": "drafted"}, ["Jane Doe"]),
            ({"channel": "email"}, ["Raj Patel"]),
            ({"source": "post"}, ["Jane Doe"]),
            ({"source": "search"}, ["Raj Patel"]),
            ({"q": "beta"}, ["Raj Patel"]),
            ({"q": "jane@acme"}, ["Jane Doe"]),
            ({"status": "drafted", "channel": "email"}, []),
            ({"q": "nobody"}, []),
        ],
    )
    def test_filters_match_the_leads_page(self, app, db, leads, params, expected):
        assert names(client_as(app, db, UserRole.VIEWER).get("/api/export.xlsx", params=params)) == expected

    @pytest.mark.parametrize("params", [{"status": "bogus"}, {"channel": "fax"}, {"source": "magic"}, {"q": "x" * 101}])
    def test_bad_filters_are_rejected(self, app, db, params):
        assert client_as(app, db, UserRole.VIEWER).get("/api/export.xlsx", params=params).status_code == 422


class TestSheetInfo:
    def test_hidden_when_not_configured(self, app, db, monkeypatch):
        monkeypatch.setattr(settings, "google_sheet_id", "")
        body = client_as(app, db, UserRole.VIEWER).get("/api/sheet").json()
        assert body == {"configured": False, "url": None, "last_synced_at": None}

    def test_link_and_last_sync(self, app, db, monkeypatch):
        monkeypatch.setattr(settings, "google_sheet_id", "abc/123")
        http = client_as(app, db, UserRole.VIEWER)
        body = http.get("/api/sheet").json()
        assert body["configured"] is True and body["last_synced_at"] is None
        assert body["url"] == "https://docs.google.com/spreadsheets/d/abc%2F123/edit"
        sheets.record_sync(db)
        assert http.get("/api/sheet").json()["last_synced_at"] is not None

    def test_garbage_last_sync_value_is_ignored(self, db):
        db.add(Setting(key=sheets.LAST_SYNC_KEY, value="not a date"))
        db.commit()
        assert sheets.last_synced(db) is None

    def test_needs_a_login(self, app, db):
        assert client_as(app, db, None).get("/api/sheet").status_code == 401


class TestSyncNow:
    @pytest.fixture(autouse=True)
    def configured(self, monkeypatch):
        monkeypatch.setattr(settings, "google_sheet_id", "sheet-id")

    def test_viewer_cannot_sync(self, app, db):
        assert client_as(app, db, UserRole.VIEWER).post("/api/sheet/sync").status_code == 403

    @pytest.mark.parametrize("role", [UserRole.MEMBER, UserRole.ADMIN])
    def test_member_and_admin_can_sync_and_it_records_the_time(self, app, db, monkeypatch, role):
        monkeypatch.setattr(sheets, "sync_sheet", AsyncMock(return_value=SyncReport(updated=2, added=1)))
        http = client_as(app, db, role)
        r = http.post("/api/sheet/sync")
        assert r.status_code == 200
        assert (r.json()["updated"], r.json()["added"]) == (2, 1) and r.json()["last_synced_at"]
        assert http.get("/api/sheet").json()["last_synced_at"] == r.json()["last_synced_at"]

    def test_failed_sync_is_a_502_and_records_nothing(self, app, db, monkeypatch):
        monkeypatch.setattr(sheets, "sync_sheet", AsyncMock(return_value=None))
        http = client_as(app, db, UserRole.MEMBER)
        assert http.post("/api/sheet/sync").status_code == 502
        assert http.get("/api/sheet").json()["last_synced_at"] is None

    def test_not_configured_is_a_409(self, app, db, monkeypatch):
        monkeypatch.setattr(settings, "google_sheet_id", "")
        assert client_as(app, db, UserRole.MEMBER).post("/api/sheet/sync").status_code == 409

    def test_needs_the_csrf_header(self, app, db):
        http = client_as(app, db, UserRole.MEMBER)
        del http.headers[security.CSRF_HEADER]
        assert http.post("/api/sheet/sync").status_code == 403

    def test_scheduler_job_records_the_time_too(self, db, monkeypatch):
        import asyncio

        from app import scheduler

        monkeypatch.setattr(sheets, "sync_sheet", AsyncMock(return_value=SyncReport()))

        async def in_session(work):
            return await work(db)

        monkeypatch.setattr(scheduler, "in_session", in_session)
        asyncio.run(scheduler.run_sheets_sync())
        assert sheets.last_synced(db) is not None
