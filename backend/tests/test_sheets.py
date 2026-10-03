from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.dialects import postgresql

from app.api import export
from app.channels import sheets
from app.channels.sheets import HEADER, SheetsClient, SheetsError, build_rows, rows_statement, sync_sheet, to_row
from app.config import settings
from app.db import get_session
from app.db.models import ChannelType, ContactRole, OutreachState

SENT = datetime(2026, 10, 5, 6, 0, tzinfo=timezone.utc)  # 11:30 IST


def make_row(outreach_id=1, **kw):
    outreach = SimpleNamespace(
        id=outreach_id, post_id=kw.get("post_id", 9), channel=kw.get("channel", ChannelType.LINKEDIN),
        state=kw.get("state", OutreachState.DM_SENT), sent_at=kw.get("sent_at", SENT),
        follow_up_due=kw.get("follow_up_due", SENT + timedelta(days=4)),
    )
    contact = SimpleNamespace(name=kw.get("person", "Jane Doe"), role=ContactRole.HR, title=kw.get("title", "Head of HR"))
    company = SimpleNamespace(name=kw.get("company", "Acme"))
    return to_row(outreach, contact, company, kw.get("post_url", "https://linkedin.com/posts/1"),
                  kw.get("last_reply", SENT + timedelta(hours=2)))


class FakeWorksheet:
    def __init__(self, column_a=None):
        self.column_a = column_a if column_a is not None else []
        self.batches: list = []
        self.appended: list = []
        self.append_kwargs: dict = {}
        self.batch_kwargs: dict = {}

    def col_values(self, col):
        assert col == 1
        return list(self.column_a)

    def batch_update(self, data, **kwargs):
        self.batches.extend(data)
        self.batch_kwargs = kwargs

    def append_rows(self, rows, **kwargs):
        self.appended.extend(rows)
        self.append_kwargs = kwargs


class TestRows:
    def test_row_columns(self):
        assert make_row(5) == [
            5, "Acme", "Jane Doe", "hr (Head of HR)", "LinkedIn post", "LinkedIn", "dm_sent",
            "2026-10-05 11:30", "2026-10-05 13:30", "2026-10-09 11:30", "https://linkedin.com/posts/1",
        ]
        assert len(make_row()) == len(HEADER)

    def test_blanks_for_missing_values(self):
        row = make_row(post_id=None, channel=ChannelType.EMAIL, sent_at=None, follow_up_due=None,
                       last_reply=None, post_url=None, title=None)
        assert row[3] == "hr" and row[4] == "Company search" and row[5] == "Email"
        assert row[7] == row[8] == row[9] == row[10] == ""

    def test_build_rows_runs_query(self):
        db = MagicMock()
        outreach = SimpleNamespace(id=1, post_id=None, channel=ChannelType.EMAIL,
                                   state=OutreachState.DRAFTED, sent_at=None, follow_up_due=None)
        contact = SimpleNamespace(name="Raj", role=ContactRole.FOUNDER, title=None)
        db.execute.return_value.all.return_value = [
            (outreach, contact, SimpleNamespace(name="Beta"), None, None)
        ]
        assert build_rows(db)[0][:3] == [1, "Beta", "Raj"]

    def test_statement_compiles_with_outer_joins(self):
        sql = str(rows_statement().compile(dialect=postgresql.dialect()))
        assert "LEFT OUTER JOIN posts" in sql and "max(messages.sent_at)" in sql
        assert "messages.direction" in sql and "ORDER BY outreaches.id" in sql


class TestUpsert:
    def test_empty_sheet_gets_header_and_appends(self):
        ws = FakeWorksheet([])
        report = SheetsClient(worksheet=ws).upsert([make_row(1), make_row(2)])
        assert ws.batches == [{"range": "A1:K1", "values": [HEADER]}]
        assert [r[0] for r in ws.appended] == [1, 2]
        assert (report.added, report.updated) == (2, 0)
        assert ws.append_kwargs["value_input_option"] == "RAW"

    def test_existing_rows_updated_in_place_new_appended(self):
        ws = FakeWorksheet(["ID", "4", "7", "9"])
        report = SheetsClient(worksheet=ws).upsert([make_row(7, person="Updated"), make_row(10)])
        assert ws.batches == [{"range": "A3:K3", "values": [make_row(7, person="Updated")]}]
        assert [r[0] for r in ws.appended] == [10]
        assert (report.updated, report.added) == (1, 1)
        assert ws.batch_kwargs["value_input_option"] == "RAW"

    def test_id_match_is_string_based(self):
        ws = FakeWorksheet(["ID", "12"])
        SheetsClient(worksheet=ws).upsert([make_row(12)])
        assert ws.batches[0]["range"] == "A2:K2" and ws.appended == []

    def test_no_writes_when_nothing_to_do(self):
        ws = FakeWorksheet(["ID"])
        report = SheetsClient(worksheet=ws).upsert([])
        assert ws.batches == [] and ws.appended == []
        assert (report.updated, report.added) == (0, 0)

    def test_header_restored_if_missing(self):
        ws = FakeWorksheet(["something", "else"])
        SheetsClient(worksheet=ws).upsert([make_row(1)])
        assert ws.batches[0] == {"range": "A1:K1", "values": [HEADER]}

    def test_blank_cells_in_id_column_are_ignored(self):
        ws = FakeWorksheet(["ID", "", "3"])
        SheetsClient(worksheet=ws).upsert([make_row(3)])
        assert ws.batches[0]["range"] == "A3:K3"


class TestClientAuth:
    def test_uses_service_account_and_sheet_id(self, monkeypatch):
        import gspread

        worksheet = FakeWorksheet()
        account = MagicMock()
        account.open_by_key.return_value.sheet1 = worksheet
        service_account = MagicMock(return_value=account)
        monkeypatch.setattr(gspread, "service_account", service_account)
        client = SheetsClient(sheet_id="SHEET123", credentials_file="/creds/sa.json")
        assert client.worksheet is worksheet
        assert service_account.call_args.kwargs["filename"] == "/creds/sa.json"
        account.open_by_key.assert_called_once_with("SHEET123")

    def test_missing_sheet_id(self):
        with pytest.raises(SheetsError):
            SheetsClient(sheet_id="").worksheet


class TestSync:
    @pytest.mark.asyncio
    async def test_sync_upserts_db_rows(self):
        ws = FakeWorksheet(["ID"])
        report = await sync_sheet(MagicMock(), SheetsClient(worksheet=ws), rows=[make_row(1)])
        assert report.added == 1 and ws.appended[0][0] == 1

    @pytest.mark.asyncio
    async def test_sync_disabled_without_sheet_id(self, monkeypatch):
        monkeypatch.setattr(settings, "google_sheet_id", "")
        assert await sync_sheet(MagicMock(), SheetsClient(sheet_id="")) is None

    @pytest.mark.asyncio
    async def test_sync_failure_is_swallowed(self):
        ws = FakeWorksheet()
        ws.col_values = MagicMock(side_effect=RuntimeError("quota exceeded"))
        assert await sync_sheet(MagicMock(), SheetsClient(worksheet=ws), rows=[make_row(1)]) is None


class TestExport:
    def client(self, monkeypatch, rows):
        monkeypatch.setattr(export, "build_rows", lambda db: rows)
        app = FastAPI()
        app.include_router(export.router)
        app.dependency_overrides[get_session] = lambda: MagicMock()
        return TestClient(app)

    def test_xlsx_download(self, monkeypatch):
        r = self.client(monkeypatch, [make_row(1), make_row(2, company="Beta")]).get("/export.xlsx")
        assert r.status_code == 200
        assert r.headers["content-type"] == export.XLSX_TYPE
        assert r.headers["content-disposition"].startswith('attachment; filename="outreach-')
        sheet = load_workbook(BytesIO(r.content)).active
        assert [c.value for c in sheet[1]] == HEADER
        assert sheet.max_row == 3
        assert [c.value for c in sheet[3]][:3] == [2, "Beta", "Jane Doe"]
        assert sheet.freeze_panes == "A2"

    def test_empty_export_has_header_only(self, monkeypatch):
        sheet = load_workbook(BytesIO(self.client(monkeypatch, []).get("/export.xlsx").content)).active
        assert sheet.max_row == 1

    def test_formula_like_text_stays_text(self, monkeypatch):
        rows = [make_row(1, person="=HYPERLINK(\"http://evil\")")]
        sheet = load_workbook(BytesIO(self.client(monkeypatch, rows).get("/export.xlsx").content)).active
        cell = sheet["C2"]
        assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
